import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import json
import pathlib

class PortalSpreadsheet:
    def __init__(self, base_url: str):
        self.base_url = base_url.rstrip("/")

    def generate_excel(self, json_payloads: dict, output_path: str, formatting_rules: dict = None):
        """Converts multiple JSON lists into distinct Excel tabs. json_payloads = {"Tab1": [dict, dict], "Tab2": [dict, dict]}"""
        wb = openpyxl.Workbook()
        wb.remove(wb.active) # Remove default sheet

        for sheet_name, data in json_payloads.items():
            ws = wb.create_sheet(title=sheet_name[:31]) # Excel limits tab names to 31 chars
            if not data:
                continue

            # 1. Write Headers
            headers = list(data[0].keys())
            ws.append(headers)
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")

            # 2. Write Data
            for row_data in data:
                row = [str(row_data.get(h, "")) if isinstance(row_data.get(h), (dict, list)) else row_data.get(h, "") for h in headers]
                ws.append(row)

            # 3. Apply Self-Organization (Auto-width and Auto-filter)
            for col_num, column_cells in enumerate(ws.columns, 1):
                length = max((len(str(cell.value)) for cell in column_cells if cell.value), default=10)
                ws.column_dimensions[get_column_letter(col_num)].width = min(length + 2, 50)
            ws.auto_filter.ref = ws.dimensions

            # 4. Apply Custom Formatting & Formulas (if provided)
            if formatting_rules and sheet_name in formatting_rules:
                rules = formatting_rules[sheet_name]
                for cell_ref, formula in rules.get("formulas", {}).items():
                    ws[cell_ref] = formula

                # Example conditional: {"column": "B", "condition": ">100", "color": "FF0000"}
                # (Can be expanded with openpyxl.formatting.rule)

        wb.save(output_path)
        return output_path

    def render_shell(self, did: str, title: str, json_data: dict) -> str:
        """UI Shell mirroring the document editor for the portal."""
        # Simple HTML table generation for the preview
        preview_html = ""
        for tab_name, rows in json_data.items():
            preview_html += f"<h4 style='color: var(--accent);'>{tab_name}</h4><table class='ui-table' style='width:100%; border-collapse:collapse; margin-bottom: 1rem;'>"
            if rows:
                preview_html += "<tr>" + "".join(f"<th style='border:1px solid var(--border); padding:0.4rem;'>{h}</th>" for h in rows[0].keys()) + "</tr>"
                for row in rows[:20]: # Limit preview to 20 rows
                    preview_html += "<tr>" + "".join(f"<td style='border:1px solid var(--border); padding:0.4rem; font-size:0.8rem;'>{str(v)[:50]}</td>" for v in row.values()) + "</tr>"
            preview_html += "</table>"

        download_url = f"{self.base_url}/download/{did}"
        return f"""
        <div id="spreadsheet-shell-{did}" class="editor-shell wide-layout">
            <div class="editor-toolbar">
                <input id="sheet-title-{did}" type="text" value="{title}" class="doc-title-input">
                <div class="toolbar-group">
                    <button class="btn-icon" title="Convert to Excel" onclick="window.location.href='{download_url}'">&#x2B07; XLSX</button>
                    <button class="btn-icon" title="Apply Formulas" hx-post="/sheet/formulas/{did}">&#x1D4A1;</button>
                </div>
            </div>
            <div class="editor-content-wrapper" style="padding: 1rem; overflow-y: auto;">
                <div style="margin-bottom: 1rem; font-size: 0.85rem; color: var(--text_muted);">Data Preview (Truncated)</div>
                {preview_html}
            </div>
        </div>
        """